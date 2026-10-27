#!/usr/bin/env python3
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_professional_excel(data, filename="output.xlsx", title=None, sheet_name="Sheet1"):
    """
    Create a professional Excel file with standardized styling
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    if title:
        ws['A1'] = title
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(data[0]) if data else 1)
        ws['A1'].alignment = Alignment(horizontal='center')
        start_row = 3
    else:
        start_row = 1
    
    # Header style
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    # Border style
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Write headers
    if data:
        headers = data[0]
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col_num)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Write data
        for row_num, row_data in enumerate(data[1:], start_row + 1):
            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.value = value
                cell.border = thin_border
                cell.alignment = Alignment(vertical='center')
        
        # Auto-adjust column widths
        for col_num in range(1, len(headers) + 1):
            column_letter = get_column_letter(col_num)
            max_length = 0
            for row in ws[column_letter]:
                try:
                    if len(str(row.value)) > max_length:
                        max_length = len(str(row.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    wb.save(filename)
    return filename

if __name__ == "__main__":
    # Example usage
    data = [
        ["Name", "Department", "Salary", "Start Date"],
        ["John Doe", "Sales", 50000, "2023-01-15"],
        ["Jane Smith", "Marketing", 55000, "2022-06-01"],
    ]
    create_professional_excel(data, "example.xlsx", title="Employee Report")
    print("Created example.xlsx")
